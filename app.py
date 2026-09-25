# -*- coding: utf-8 -*-
"""
Suite Financiera — For Drink SA
====================================================================
Módulos:
1. 💵 Gestión de Tesorería (Planes de Pago y Escaleras ARS)
2. 🏦 Clearing / Cheques Emitidos (Cámaras, Calendario y Anulaciones)
3. 🔍 Analizador de Libradores · BCRA (Scoring Nativo)
"""

from __future__ import annotations

import calendar
import datetime as dt
import importlib
import re
import time
from io import BytesIO
from pathlib import Path

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    from streamlit_gsheets import GSheetsConnection
    HAS_GSHEETS = True
except ImportError:
    HAS_GSHEETS = False

import core
importlib.reload(core)

# ---------------------------------------------------------------------------
# Configuración general y estilos
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="Finanzas — For Drink SA",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&family=JetBrains+Mono:wght@500;600;700&display=swap');

    html, body, [class*="css"], .stMarkdown {
        font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif !important;
    }

    [data-testid="stMetricValue"] {
        font-family: 'JetBrains Mono', monospace !important;
        font-size: 1.45rem !important;
        font-weight: 700 !important;
        color: #f8fafc !important;
        white-space: nowrap !important;
    }
    [data-testid="stMetricLabel"] {
        font-size: 0.8rem !important;
        font-weight: 600 !important;
        text-transform: uppercase;
        letter-spacing: 0.05em;
        color: #94a3b8 !important;
    }

    h1 {
        font-size: 1.65rem !important;
        font-weight: 700 !important;
        letter-spacing: -0.025em !important;
        color: #f8fafc !important;
        margin-bottom: 0.15rem !important;
    }
    h2, h3 {
        font-size: 1.25rem !important;
        font-weight: 600 !important;
        letter-spacing: -0.015em !important;
        color: #e2e8f0 !important;
    }

    /* Tarjetas del Calendario Original */
    .cal-day-box {
        background: #18202a;
        border: 1px solid #2d3748;
        border-radius: 8px;
        padding: 8px 10px;
        min-height: 96px;
        display: flex;
        flex-direction: column;
        justify-content: space-between;
    }
    .cal-day-header {
        display: flex;
        justify-content: space-between;
        align-items: center;
        font-size: 12px;
        font-weight: 700;
        color: #94a3b8;
    }
    .cal-day-total {
        font-family: 'JetBrains Mono', monospace !important;
        font-size: 13px;
        font-weight: 800;
        color: #38bdf8;
        margin-top: 4px;
    }
    .cal-details-summary {
        font-size: 10px;
        color: #94a3b8;
        cursor: pointer;
        outline: none;
        user-select: none;
        margin-top: 4px;
    }
    .cal-details-summary:hover {
        color: #38bdf8;
    }

    /* Badges de bancos */
    .bank-pill {
        display: inline-flex;
        align-items: center;
        gap: 6px;
        padding: 4px 10px;
        border-radius: 9999px;
        font-size: 11.5px;
        font-weight: 600;
        margin-right: 6px;
        margin-bottom: 6px;
    }
    .bank-dot {
        width: 8px;
        height: 8px;
        border-radius: 50%;
        display: inline-block;
    }

    /* Tarjetas BCRA */
    .bcra-card-dark {
        background: #18202a;
        border: 1px solid #2d3748;
        border-radius: 12px;
        padding: 16px 20px;
    }
    .bcra-label-dark {
        font-size: 11px;
        text-transform: uppercase;
        font-weight: 700;
        letter-spacing: 0.05em;
        color: #94a3b8;
    }
    .bcra-val-dark {
        font-family: 'JetBrains Mono', monospace !important;
        font-size: 26px;
        font-weight: 800;
        margin-top: 4px;
        line-height: 1.1;
    }
    .bcra-table-box {
        background: #18202a;
        border: 1px solid #2d3748;
        border-radius: 12px;
        padding: 6px 12px;
    }
    .bcra-drawer-box {
        background: #18202a;
        border: 1px solid #2d3748;
        border-radius: 12px;
        padding: 18px 20px;
    }
    .kpi-mini-box {
        background: #1e2634;
        border: 1px solid #334155;
        border-radius: 8px;
        padding: 10px 12px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

SEED_PATH = Path(__file__).parent / "data_seed.csv"

# ---------------------------------------------------------------------------
# Control de Acceso
# ---------------------------------------------------------------------------

params = st.query_params
is_admin_param = params.get("admin", "") == "1"

if "is_admin" not in st.session_state:
    st.session_state.is_admin = is_admin_param

# ---------------------------------------------------------------------------
# Conexión persistente a Google Sheets
# ---------------------------------------------------------------------------

conn = None
if HAS_GSHEETS:
    try:
        conn = st.connection("gsheets", type=GSheetsConnection)
    except Exception:
        conn = None


def load_data_source() -> pd.DataFrame:
    if conn is not None:
        try:
            df_cloud = conn.read(ttl="0s")
            if df_cloud is not None and not df_cloud.empty:
                return core.normalize_df(df_cloud)
        except Exception as e:
            st.sidebar.warning(f"Aviso Planes de Pago: {e}")
    if SEED_PATH.exists():
        return core.load_seed_csv(SEED_PATH)
    return core.normalize_df(pd.DataFrame(columns=core.COLUMNS))


def save_data_source(df: pd.DataFrame) -> bool:
    if not st.session_state.is_admin:
        st.error("Acceso de solo lectura: no tiene permisos para modificar la base.")
        return False

    df_to_save = core.normalize_df(df)
    df_out = df_to_save[core.COLUMNS].copy()
    if "fecha" in df_out.columns:
        df_out["fecha"] = df_out["fecha"].dt.strftime("%Y-%m-%d")

    if conn is not None:
        try:
            conn.update(data=df_out)
            return True
        except Exception as e:
            st.error(f"Error al sincronizar con Google Sheets: {e}")
            return False
    return False


def load_cheques_source() -> pd.DataFrame:
    if conn is not None:
        try:
            df_ch = conn.read(worksheet="Cheques_Emitidos", ttl="0s")
            if df_ch is not None and not df_ch.empty:
                return core.normalize_cheques_df(df_ch)
        except Exception as e:
            st.sidebar.warning(f"Aviso Cheques Emitidos: {e}")
    return core.normalize_cheques_df(pd.DataFrame())


def save_cheques_source(df: pd.DataFrame) -> bool:
    if not st.session_state.is_admin:
        st.error("Acceso de solo lectura.")
        return False

    df_out = core.prepare_cheques_to_save(df)
    if conn is not None:
        try:
            conn.update(worksheet="Cheques_Emitidos", data=df_out)
            return True
        except Exception as e:
            st.error(f"Error al guardar cheques en Google Sheets: {e}")
            return False
    return False


def init_state() -> None:
    if "df" not in st.session_state:
        st.session_state.df = load_data_source()
    if "df_cheques" not in st.session_state:
        st.session_state.df_cheques = load_cheques_source()
    if "editing_id" not in st.session_state:
        st.session_state.editing_id = None
    if "sede_global" not in st.session_state:
        st.session_state.sede_global = core.SEDE_CONSOLIDADO
    if "feriados" not in st.session_state:
        st.session_state.feriados = []
    if "semaforos_mes" not in st.session_state:
        st.session_state.semaforos_mes = {}


init_state()


def get_df() -> pd.DataFrame:
    return st.session_state.df


def set_df(df: pd.DataFrame, sync_cloud: bool = True) -> None:
    norm = core.normalize_df(df)
    st.session_state.df = norm
    if sync_cloud and st.session_state.is_admin:
        save_data_source(norm)


def get_cheques() -> pd.DataFrame:
    return st.session_state.df_cheques


def set_cheques(df: pd.DataFrame, sync_cloud: bool = True) -> None:
    norm = core.normalize_cheques_df(df)
    st.session_state.df_cheques = norm
    if sync_cloud and st.session_state.is_admin:
        save_cheques_source(norm)


def fmt_ars(value: float) -> str:
    try:
        return f"$ {value:,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return "$ 0"


BANK_THEMES = {
    "GALICIA": {"color": "#ff7a00", "bg": "rgba(255, 122, 0, 0.12)", "border": "rgba(255, 122, 0, 0.35)", "label": "Galicia"},
    "MACRO": {"color": "#0284c7", "bg": "rgba(2, 132, 199, 0.12)", "border": "rgba(2, 132, 199, 0.35)", "label": "Macro"},
    "SANTANDER": {"color": "#ef4444", "bg": "rgba(239, 68, 68, 0.12)", "border": "rgba(239, 68, 68, 0.35)", "label": "Santander"},
    "BBVA": {"color": "#3b82f6", "bg": "rgba(59, 130, 246, 0.12)", "border": "rgba(59, 130, 246, 0.35)", "label": "BBVA"},
    "BNA": {"color": "#38bdf8", "bg": "rgba(56, 189, 248, 0.12)", "border": "rgba(56, 189, 248, 0.35)", "label": "BNA Nación"},
    "SUPERVIELLE": {"color": "#e11d48", "bg": "rgba(225, 29, 72, 0.12)", "border": "rgba(225, 29, 72, 0.35)", "label": "Supervielle"},
    "PROVINCIA": {"color": "#10b981", "bg": "rgba(16, 185, 129, 0.12)", "border": "rgba(16, 185, 129, 0.35)", "label": "BAPRO"},
}


# ---------------------------------------------------------------------------
# Barra lateral: Finanzas (For Drink SA)
# ---------------------------------------------------------------------------

st.sidebar.title("💼 Finanzas")
st.sidebar.caption("For Drink SA — Control Financiero")

if conn is not None:
    st.sidebar.markdown(
        """
        <div style="display:flex; align-items:center; gap:6px; color:#10b981; font-size:0.8rem; font-weight:600; margin-bottom:12px;">
            <span style="height:7px; width:7px; background-color:#10b981; border-radius:50%; display:inline-block;"></span>
            Google Sheets Conectado
        </div>
        """,
        unsafe_allow_html=True,
    )

if st.session_state.is_admin:
    modulo_activo = st.sidebar.radio(
        "Módulo Activo",
        [
            "💵 Gestión de Tesorería",
            "🏦 Clearing / Cheques Emitidos",
            "🔍 Analizador de Libradores · BCRA"
        ],
        index=0,
    )
else:
    modulo_activo = "💵 Gestión de Tesorería"

st.sidebar.divider()

if not st.session_state.is_admin:
    with st.sidebar.expander("🔒 Acceso Admin", expanded=False):
        admin_pass = st.text_input("Clave de edición", type="password", key="pass_admin_input")
        if st.button("Habilitar"):
            if admin_pass == "fordrink2026":
                st.session_state.is_admin = True
                st.rerun()
            else:
                st.error("Clave incorrecta")
else:
    st.sidebar.markdown("<small style='color:#38bdf8; font-weight:600;'>🔑 Administrador Habilitado</small>", unsafe_allow_html=True)
    if st.sidebar.button("Cerrar sesión Admin"):
        st.session_state.is_admin = False
        st.rerun()

# ===========================================================================
# MÓDULO 1: GESTIÓN DE TESORERÍA
# ===========================================================================
if modulo_activo == "💵 Gestión de Tesorería":
    st.sidebar.markdown("##### 🏢 Sede / Tesorería")
    sede_global = st.sidebar.segmented_control(
        "Sede",
        options=core.OPCIONES_SEDE,
        default=st.session_state.sede_global,
        label_visibility="collapsed",
    )
    if sede_global is None:
        sede_global = core.SEDE_CONSOLIDADO
    st.session_state.sede_global = sede_global

    umbral_tension = 30_000_000.0 if sede_global == core.SEDE_CONSOLIDADO else 15_000_000.0

    if st.session_state.is_admin:
        with st.sidebar.expander("📁 Importar / Actualizar base", expanded=False):
            uploaded = st.file_uploader("Subir archivo Excel", type=["xlsx"], key="uploader_replace")
            if uploaded is not None and st.button("Reemplazar base actual", type="primary"):
                try:
                    new_df = core.load_from_uploaded_xlsx(uploaded)
                    set_df(new_df, sync_cloud=True)
                    st.success(f"Cargados {len(new_df)} pagos.")
                    st.rerun()
                except Exception as e:
                    st.error(f"Error: {e}")

            st.markdown("---")
            if st.button("🔄 Recargar datos desde Google Sheets"):
                st.session_state.df = load_data_source()
                st.rerun()
        st.sidebar.divider()

    st.sidebar.subheader("🔎 Filtros")
    df_all_raw = get_df()
    df_all = core.filter_by_sede(df_all_raw, sede_global)

    proveedores_disponibles = sorted(df_all["proveedor"].dropna().unique().tolist())
    f_proveedores = st.sidebar.multiselect("Proveedor", proveedores_disponibles, default=proveedores_disponibles)
    f_estados = st.sidebar.multiselect("Estado", core.ESTADOS_VALIDOS, default=core.ESTADOS_VALIDOS)

    min_date = df_all["fecha"].min()
    max_date = df_all["fecha"].max()
    if pd.isna(min_date) or pd.isna(max_date):
        min_date = max_date = pd.Timestamp.today()

    f_rango = st.sidebar.date_input(
        "Rango de fechas",
        value=(min_date.date(), max_date.date()),
        min_value=min_date.date(),
        max_value=max_date.date(),
    )

    meses_disponibles = sorted(df_all["mes"].dropna().unique().tolist())
    f_meses = st.sidebar.multiselect("Mes", meses_disponibles, default=meses_disponibles)

    def apply_filters(df: pd.DataFrame) -> pd.DataFrame:
        out = df.copy()
        if f_proveedores:
            out = out[out["proveedor"].isin(f_proveedores)]
        if f_estados:
            out = out[out["estado"].isin(f_estados)]
        if f_meses:
            out = out[out["mes"].isin(f_meses)]
        if isinstance(f_rango, tuple) and len(f_rango) == 2:
            d0, d1 = pd.Timestamp(f_rango[0]), pd.Timestamp(f_rango[1])
            out = out[(out["fecha"] >= d0) & (out["fecha"] <= d1)]
        return out

    df_filtered = apply_filters(df_all)

    st.title("Gestión de Tesorería — Planes de Pago")
    st.caption(f"For Drink SA · Sede: **{sede_global}** · Control de caja proyectado.")

    if st.session_state.is_admin:
        tab_kpi, tab_matriz, tab_escaleras, tab_conciliar, tab_asistente, tab_abm, tab_export = st.tabs([
            "📊 Panel Ejecutivo",
            "🗓️ Matriz Semanal",
            "🪜 Detalle de Escaleras",
            "⚡ Conciliación Rápida",
            "➕ Cargar Nueva Escalera",
            "📝 Registro Manual (ABM)",
            "⬇️ Exportar",
        ])
    else:
        tab_kpi, tab_matriz, tab_escaleras = st.tabs([
            "📊 Panel Ejecutivo",
            "🗓️ Matriz Semanal",
            "🪜 Detalle de Escaleras",
        ])

    with tab_kpi:
        kpis = core.compute_kpis(df_filtered)
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Total Comprometido", fmt_ars(kpis.total_comprometido))
        c2.metric("Total Pagado", fmt_ars(kpis.total_pagado))
        c3.metric("Saldo Pendiente", fmt_ars(kpis.saldo_pendiente))
        c4.metric("Desembolso Promedio", fmt_ars(kpis.desembolso_promedio))

        st.divider()
        left, right = st.columns([3, 2])
        with left:
            st.subheader("Distribución mensual proyectada")
            dist = core.monthly_distribution(df_filtered)
            if not dist.empty:
                fig = px.bar(
                    dist, x="mes", y="importe_ars", text_auto=".2s",
                    labels={"mes": "Mes", "importe_ars": "Importe ARS"},
                    color_discrete_sequence=["#38bdf8"],
                )
                fig.update_layout(yaxis_title="Importe ARS", xaxis_title="", margin=dict(t=10, b=10), height=340, paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
                st.plotly_chart(fig, use_container_width=True)
            else:
                st.info("Sin datos para graficar.")

            if sede_global == core.SEDE_CONSOLIDADO and not df_filtered.empty:
                st.subheader("Distribución por Sede")
                por_sede = df_filtered.groupby("tesoreria")["importe_ars"].sum().reset_index()
                fig_sede = px.pie(por_sede, names="tesoreria", values="importe_ars", hole=0.5,
                                  color_discrete_sequence=["#0284c7", "#38bdf8"])
                fig_sede.update_layout(margin=dict(t=10, b=10), height=280, paper_bgcolor="rgba(0,0,0,0)")
                st.plotly_chart(fig_sede, use_container_width=True)

        with right:
            st.subheader("🚦 Carga Semanal Proyectada")
            st.caption(f"Semáforo rojo (> {fmt_ars(umbral_tension)}) según sede ({sede_global}).")

            hoy_ts = pd.Timestamp(dt.date.today())
            lunes_actual = hoy_ts - pd.Timedelta(days=hoy_ts.weekday())
            df_futuro = df_filtered[df_filtered["lunes_semana"] >= lunes_actual].copy()

            if df_futuro.empty:
                st.info("No hay desembolsos pendientes programados desde la semana actual.")
            else:
                sem_futuras = (
                    df_futuro.groupby(["lunes_semana", "semana_etiqueta"])["importe_ars"]
                    .sum()
                    .reset_index()
                    .sort_values("lunes_semana")
                )
                promedio_futuro = sem_futuras["importe_ars"].mean()
                st.markdown(f"**Promedio semanal:** `{fmt_ars(promedio_futuro)}` ({len(sem_futuras)} semanas)")

                for _, r_sem in sem_futuras.iterrows():
                    monto_s = r_sem["importe_ars"]
                    sem_label = r_sem["semana_etiqueta"]
                    if monto_s > umbral_tension or monto_s > (promedio_futuro * 1.3):
                        st.error(f"**{sem_label}** — {fmt_ars(monto_s)} *(Tensión Alta)*", icon="🔴")
                    elif monto_s >= (promedio_futuro * 0.8):
                        st.warning(f"**{sem_label}** — {fmt_ars(monto_s)} *(En Promedio)*", icon="🟡")
                    else:
                        st.success(f"**{sem_label}** — {fmt_ars(monto_s)} *(Holgada)*", icon="🟢")

        st.divider()
        st.subheader("Compromiso y Avance por Proveedor")
        ctrl = core.provider_control(df_filtered)
        if not ctrl.empty:
            fig2 = go.Figure()
            fig2.add_trace(go.Bar(y=ctrl["proveedor"], x=ctrl["compromiso_total_ars"], orientation="h",
                                   name="Comprometido", marker_color="#334155"))
            fig2.add_trace(go.Bar(y=ctrl["proveedor"], x=ctrl["pagado_ars"], orientation="h",
                                   name="Pagado", marker_color="#38bdf8"))
            fig2.update_layout(barmode="overlay", xaxis_title="Importe ARS", height=max(320, 26 * len(ctrl)),
                                margin=dict(t=10, b=10), legend=dict(orientation="h", yanchor="bottom", y=1.02),
                                paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
            st.plotly_chart(fig2, use_container_width=True)

    with tab_matriz:
        st.subheader("Matriz Semanal de Flujo de Efectivo")
        if df_filtered.empty:
            st.info("Sin datos para mostrar con los filtros aplicados.")
        else:
            df_sorted_weeks = df_filtered.sort_values("lunes_semana")
            semanas_ordenadas = (
                df_sorted_weeks[["lunes_semana", "semana_etiqueta"]]
                .drop_duplicates()
                .dropna()
            )

            opciones_semanas = semanas_ordenadas["semana_etiqueta"].tolist()
            hoy_ts = pd.Timestamp(dt.date.today())
            lunes_actual = hoy_ts - pd.Timedelta(days=hoy_ts.weekday())

            idx_default = 0
            semanas_desde_hoy = semanas_ordenadas[semanas_ordenadas["lunes_semana"] >= lunes_actual]
            if not semanas_desde_hoy.empty:
                primera_sem_label = semanas_desde_hoy["semana_etiqueta"].iloc[0]
                if primera_sem_label in opciones_semanas:
                    idx_default = opciones_semanas.index(primera_sem_label)

            col_filtro_sem, col_btn_excel, col_btn_pdf = st.columns([2, 1.2, 1.2])
            with col_filtro_sem:
                semana_inicio_sel = st.selectbox(
                    "📅 Mostrar semanas desde:",
                    options=opciones_semanas,
                    index=idx_default,
                    help="Abre por defecto en la semana en curso.",
                )

            lunes_corte = semanas_ordenadas.loc[
                semanas_ordenadas["semana_etiqueta"] == semana_inicio_sel, "lunes_semana"
            ].iloc[0]

            df_matriz_periodo = df_filtered[df_filtered["lunes_semana"] >= lunes_corte]
            matrix = core.build_weekly_matrix(df_matriz_periodo)

            if matrix.empty:
                st.info("No hay pagos para el período seleccionado.")
            else:
                def export_matrix_to_excel(matrix_df: pd.DataFrame) -> bytes:
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "Matriz Semanal"
                    ws["A1"] = f"MATRIZ SEMANAL DE FLUJO DE EFECTIVO — {sede_global.upper()}"
                    ws["A1"].font = Font(size=13, bold=True, color="1F4E78")
                    ws["A2"] = f"Período: Desde {semana_inicio_sel} · Generado el: {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}"
                    ws["A2"].font = Font(size=9, italic=True, color="666666")

                    headers = list(matrix_df.columns)
                    h_row = 4
                    for c_idx, h in enumerate(headers, 1):
                        cell = ws.cell(row=h_row, column=c_idx, value=h)
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                    peak_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    total_fill = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")

                    for r_idx, (_, row) in enumerate(matrix_df.iterrows()):
                        row_num = h_row + 1 + r_idx
                        is_total_row = (row.get("Semana") == "TOTAL POR PROVEEDOR")
                        total_semanal = row.get("TOTAL SEMANAL", 0)
                        is_peak = (not is_total_row and isinstance(total_semanal, (int, float)) and total_semanal > umbral_tension)

                        for c_idx, h in enumerate(headers, 1):
                            val = row[h]
                            cell = ws.cell(row=row_num, column=c_idx)

                            if is_total_row:
                                cell.font = Font(bold=True)
                                cell.fill = total_fill
                            elif is_peak and h == "TOTAL SEMANAL":
                                cell.font = Font(bold=True)
                                cell.fill = peak_fill

                            if h == "Semana":
                                cell.value = str(val)
                                cell.alignment = Alignment(horizontal="center")
                            else:
                                cell.value = float(val) if pd.notna(val) else 0.0
                                cell.number_format = "$ #,##0"

                    for col in ws.columns:
                        max_len = max(len(str(cell.value or "")) for cell in col)
                        col_letter = get_column_letter(col[0].column)
                        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

                    buf = BytesIO()
                    wb.save(buf)
                    return buf.getvalue()

                def export_matrix_to_pdf(matrix_df: pd.DataFrame) -> bytes:
                    buf = BytesIO()
                    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=25, bottomMargin=25)
                    styles = getSampleStyleSheet()
                    story = []

                    title_style = ParagraphStyle('MatTitle', parent=styles['Heading1'], fontSize=13, textColor=colors.HexColor('#1F4E78'), spaceAfter=3)
                    sub_style = ParagraphStyle('MatSub', parent=styles['Normal'], fontSize=8, textColor=colors.HexColor('#555555'), spaceAfter=10)

                    story.append(Paragraph(f"<b>FOR DRINK SA — MATRIZ SEMANAL ({sede_global.upper()})</b>", title_style))
                    story.append(Paragraph(f"Desde {semana_inicio_sel} | Emitido: {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}", sub_style))

                    headers = list(matrix_df.columns)
                    table_data = [headers]
                    for _, r in matrix_df.iterrows():
                        row_vals = []
                        for h in headers:
                            v = r[h]
                            row_vals.append(str(v) if h == "Semana" else (fmt_ars(v) if pd.notna(v) else "$ 0"))
                        table_data.append(row_vals)

                    col_w = max(40, int(780 / len(headers)))
                    t = Table(table_data, colWidths=[col_w] * len(headers))
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, 0), 7),
                        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#DDDDDD')),
                        ('FONTSIZE', (0, 1), (-1, -1), 6.5),
                        ('TOPPADDING', (0, 0), (-1, -1), 3),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#EFEFEF')),
                        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                    ]))
                    story.append(t)
                    doc.build(story)
                    return buf.getvalue()

                with col_btn_excel:
                    st.write("")
                    st.write("")
                    st.download_button(
                        label="📥 Descargar Excel",
                        data=export_matrix_to_excel(matrix),
                        file_name=f"Matriz_Semanal_{sede_global}_{dt.date.today().strftime('%Y%m%d')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )

                with col_btn_pdf:
                    st.write("")
                    st.write("")
                    st.download_button(
                        label="📄 Descargar PDF",
                        data=export_matrix_to_pdf(matrix),
                        file_name=f"Matriz_Semanal_{sede_global}_{dt.date.today().strftime('%Y%m%d')}.pdf",
                        mime="application/pdf",
                        use_container_width=True,
                    )

                st.markdown("---")
                provider_cols = [c for c in matrix.columns if c not in ("Semana", "TOTAL SEMANAL", "ACUMULADO")]
                def highlight_peaks(row):
                    styles = [""] * len(row)
                    if row["Semana"] == "TOTAL POR PROVEEDOR":
                        return ["font-weight: 700; background-color: #1e293b; color: #38bdf8;"] * len(row)
                    try:
                        if row["TOTAL SEMANAL"] > umbral_tension:
                            idx = list(row.index).index("TOTAL SEMANAL")
                            styles[idx] = "background-color: #7f1d1d; font-weight: 700; color: #fecaca;"
                    except Exception:
                        pass
                    return styles

                styled = matrix.style.apply(highlight_peaks, axis=1).format({c: fmt_ars for c in provider_cols + ["TOTAL SEMANAL", "ACUMULADO"]})
                st.dataframe(styled, use_container_width=True, height=520, hide_index=True)

    with tab_escaleras:
        st.subheader("🪜 Control Individual de Escaleras")
        prov_list = sorted(df_filtered["proveedor"].dropna().unique().tolist())
        if not prov_list:
            st.info("No hay proveedores para los filtros seleccionados.")
        else:
            col_sel_p, col_sel_op = st.columns([1, 1])
            with col_sel_p:
                p_sel = st.selectbox("Seleccionar Proveedor", prov_list, key="p_inspect")

            df_p_raw = df_filtered[df_filtered["proveedor"] == p_sel].copy()

            def clean_op_label(r):
                c = str(r["concepto"]).strip() if pd.notna(r["concepto"]) and str(r["concepto"]).strip().lower() != "nan" else ""
                return c if c else "Operación General / Presupuesto Base"

            df_p_raw["operacion_identificada"] = df_p_raw.apply(clean_op_label, axis=1)
            operaciones_encontradas = sorted(df_p_raw["operacion_identificada"].unique().tolist())

            with col_sel_op:
                if len(operaciones_encontradas) > 1:
                    opciones_operacion = ["Todas las Operaciones (Consolidado)"] + operaciones_encontradas
                    op_sel = st.selectbox(f"📌 Operación ({len(operaciones_encontradas)} registradas)", opciones_operacion, key="op_inspect")
                else:
                    op_sel = "Todas las Operaciones (Consolidado)"
                    st.selectbox("📌 Operación", ["Operación Única"], disabled=True, key="op_inspect_single")

            df_p = df_p_raw[df_p_raw["operacion_identificada"] == op_sel].copy() if op_sel != "Todas las Operaciones (Consolidado)" else df_p_raw.copy()
            subtitulo_op = f" — {op_sel}" if op_sel != "Todas las Operaciones (Consolidado)" else ""

            kpis_p = core.compute_kpis(df_p)
            st.markdown(f"#### Ficha: **{p_sel}**{subtitulo_op}")
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Compromiso Total", fmt_ars(kpis_p.total_comprometido))
            m2.metric("Total Pagado", fmt_ars(kpis_p.total_pagado))
            m3.metric("Saldo Pendiente", fmt_ars(kpis_p.saldo_pendiente))
            pct_p = (kpis_p.total_pagado / kpis_p.total_comprometido * 100) if kpis_p.total_comprometido else 0
            m4.metric("% Cancelado", f"{pct_p:.1f}%")
            st.progress(min(pct_p / 100, 1.0))

            st.divider()

            todas_obs = [str(o).strip() for o in df_p["obs"].dropna().unique() if str(o).strip() and str(o).strip().lower() != "nan"]
            with st.expander("🔍 Observaciones del acuerdo"):
                if todas_obs:
                    for idx_o, obs_t in enumerate(todas_obs, 1):
                        st.markdown(f"**Nota {idx_o}:** {obs_t}")
                else:
                    st.caption("Sin notas adicionales registradas.")

            def merge_concepto_obs(r):
                c = str(r["concepto"]).strip() if pd.notna(r["concepto"]) and str(r["concepto"]).strip().lower() != "nan" else ""
                o = str(r["obs"]).strip() if pd.notna(r["obs"]) and str(r["obs"]).strip().lower() != "nan" else ""
                if c and o:
                    return f"{c} ({o})" if c != o else c
                return c or o or "Presupuesto Base"

            df_p["presupuesto_obs"] = df_p.apply(merge_concepto_obs, axis=1)

            def format_ladder_table(sub_df: pd.DataFrame) -> pd.DataFrame:
                t = sub_df.sort_values("fecha").copy()
                return pd.DataFrame({
                    "Fecha Vto.": t["fecha"],
                    "Importe ARS": t["importe_ars"],
                    "Estado": t["estado"],
                    "Detalle / Observaciones": t["presupuesto_obs"],
                })

            def export_provider_to_excel(df_table: pd.DataFrame, prov_name: str, op_name: str) -> bytes:
                wb = Workbook()
                ws = wb.active
                ws.title = "Escalera"
                ws["A1"] = f"ESCALERA DE PAGOS — {prov_name.upper()}" + (f" ({op_name.upper()})" if op_name != "Todas las Operaciones (Consolidado)" else "")
                ws["A1"].font = Font(size=13, bold=True, color="1F4E78")
                ws["A2"] = f"Generado el: {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}"
                ws["A2"].font = Font(size=9, italic=True, color="666666")

                headers = list(df_table.columns)
                h_row = 4
                for col_num, header in enumerate(headers, 1):
                    cell = ws.cell(row=h_row, column=col_num, value=header)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                total_suma = 0.0
                for r_idx, (_, row) in enumerate(df_table.iterrows()):
                    row_num = h_row + 1 + r_idx
                    for c_idx, h in enumerate(headers, 1):
                        val = row[h]
                        cell = ws.cell(row=row_num, column=c_idx)
                        if h == "Fecha Vto." and pd.notna(val):
                            cell.value = val.strftime("%d/%m/%Y") if hasattr(val, "strftime") else str(val)
                            cell.alignment = Alignment(horizontal="center")
                        elif h == "Importe ARS":
                            v_float = float(val) if pd.notna(val) else 0.0
                            total_suma += v_float
                            cell.value = v_float
                            cell.number_format = "$ #,##0"
                        else:
                            cell.value = str(val) if pd.notna(val) else ""

                tot_row = h_row + 1 + len(df_table)
                ws.cell(row=tot_row, column=1, value="TOTAL").font = Font(bold=True)
                ws.cell(row=tot_row, column=1).alignment = Alignment(horizontal="center")
                cell_tot = ws.cell(row=tot_row, column=2, value=total_suma)
                cell_tot.font = Font(bold=True)
                cell_tot.number_format = "$ #,##0"
                for c_idx in range(1, len(headers) + 1):
                    ws.cell(row=tot_row, column=c_idx).fill = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")

                for col in ws.columns:
                    max_len = max(len(str(cell.value or "")) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

                buf = BytesIO()
                wb.save(buf)
                return buf.getvalue()

            table_view = format_ladder_table(df_p)
            c_down_xl, _ = st.columns([1.5, 3])
            with c_down_xl:
                st.download_button("📥 Descargar Excel del Proveedor", export_provider_to_excel(table_view, p_sel, op_sel), f"Escalera_{p_sel.replace(' ', '_')}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"dl_xl_{p_sel}", use_container_width=True)

            st.dataframe(
                table_view.style.format({
                    "Fecha Vto.": lambda d: d.strftime("%d/%m/%Y") if pd.notna(d) else "",
                    "Importe ARS": fmt_ars,
                }),
                use_container_width=True,
                hide_index=True,
            )

    if st.session_state.is_admin:
        with tab_conciliar:
            st.subheader("⚡ Conciliación Rápida")
            df_prog = df_filtered[df_filtered["estado"] != "Pagado"].sort_values("fecha").copy()
            if df_prog.empty:
                st.success("🎉 No hay pagos pendientes en este rango.")
            else:
                st.write(f"Hay **{len(df_prog)} pagos pendientes**.")
                df_editor = df_prog[["id", "tesoreria", "fecha", "proveedor", "concepto", "importe_ars", "estado", "obs"]].copy()
                df_editor.insert(0, "Conciliar", False)
                df_editor["fecha"] = df_editor["fecha"].dt.date

                edited = st.data_editor(
                    df_editor,
                    column_config={
                        "Conciliar": st.column_config.CheckboxColumn("¿Pagar?", default=False),
                        "importe_ars": st.column_config.NumberColumn("Importe ARS", format="$ %d"),
                        "fecha": st.column_config.DateColumn("Fecha Vto.", format="DD/MM/YYYY"),
                    },
                    disabled=["id", "tesoreria", "fecha", "proveedor", "concepto", "importe_ars", "estado", "obs"],
                    use_container_width=True,
                    hide_index=True,
                    height=360,
                    key="conciliar_editor",
                )

                pagos_a_conciliar = edited[edited["Conciliar"] == True]["id"].tolist()
                if st.button("💾 Marcar Seleccionados como PAGADO", disabled=len(pagos_a_conciliar) == 0):
                    df_current = get_df()
                    df_updated = core.update_payments_status(df_current, pagos_a_conciliar, new_status="Pagado")
                    set_df(df_updated, sync_cloud=True)
                    st.success(f"Se actualizaron {len(pagos_a_conciliar)} pagos a estado PAGADO.")
                    st.rerun()

        with tab_asistente:
            st.subheader("➕ Cargar Nueva Escalera de Pago")
            c_cab1, c_cab2 = st.columns(2)
            with c_cab1:
                esc_tesoreria = st.selectbox("Tesorería", core.TESORERIAS_VALIDAS, index=0 if sede_global == "Tucumán" else 1, key="esc_sede")
                tipo_prov = st.radio("Proveedor", ["Existente", "Nuevo"], horizontal=True, key="esc_tipo_prov")
                prov_opts = sorted(get_df()["proveedor"].dropna().unique().tolist())
                esc_proveedor = st.selectbox("Seleccionar", prov_opts, key="esc_prov_sel") if tipo_prov == "Existente" and prov_opts else st.text_input("Nombre", key="esc_prov_new").strip().upper()

            with c_cab2:
                esc_proyecto = st.text_input("Proyecto / Obra", value="L2 TUC" if esc_tesoreria == "Tucumán" else "L2 BA", key="esc_proy")
                esc_concepto = st.text_input("Presupuesto / Concepto", placeholder="Ej: Carpintería - Presupuesto N° 450", key="esc_conc")

            st.markdown("---")
            col_nc, _ = st.columns([1, 3])
            with col_nc:
                cant_cuotas = st.number_input("Cantidad de cuotas", min_value=1, max_value=50, value=4, step=1, key="esc_n_cuotas")

            if ("grid_ladder_data" not in st.session_state 
                or len(st.session_state.grid_ladder_data) != cant_cuotas 
                or st.session_state.get("reset_grid_flag", False)):
                base_f = dt.date.today()
                st.session_state.grid_ladder_data = pd.DataFrame([
                    {"Tramo": f"Cuota {i + 1}", "Fecha Vto.": base_f + dt.timedelta(days=7 * i), "Importe ARS": 0.0}
                    for i in range(int(cant_cuotas))
                ])
                st.session_state["reset_grid_flag"] = False

            grid_edited = st.data_editor(
                st.session_state.grid_ladder_data,
                column_config={
                    "Tramo": st.column_config.TextColumn("Tramo", disabled=True),
                    "Fecha Vto.": st.column_config.DateColumn("Fecha Vto.", format="DD/MM/YYYY", required=True),
                    "Importe ARS": st.column_config.NumberColumn("Importe en $ ARS", min_value=0.0, format="$ %d", required=True),
                },
                use_container_width=True,
                hide_index=True,
                key="grid_cuotas_editor",
            )

            total_acordado = float(grid_edited["Importe ARS"].sum())
            st.info(f"📊 **Resumen:** {len(grid_edited)} tramos · Total Acordado: **{fmt_ars(total_acordado)}**")

            if st.button("🚀 Guardar Escalera en Google Sheets", use_container_width=True):
                if not esc_proveedor or not esc_concepto or total_acordado <= 0:
                    st.error("Completá proveedor, presupuesto y cuotas con importe.")
                else:
                    items_ladder = [{"fecha": pd.to_datetime(r["Fecha Vto."]), "importe": float(r["Importe ARS"])} for _, r in grid_edited.iterrows()]
                    df_curr = get_df()
                    new_rows = core.generate_custom_ladder(
                        df_existing=df_curr,
                        tesoreria=esc_tesoreria,
                        proveedor=esc_proveedor,
                        proyecto=esc_proyecto,
                        concepto=esc_concepto,
                        moneda="ARS",
                        tc=1.0,
                        items=items_ladder,
                    )
                    df_total = pd.concat([df_curr[core.COLUMNS], new_rows[core.COLUMNS]], ignore_index=True)
                    if save_data_source(df_total):
                        st.session_state.df = core.normalize_df(df_total)
                        st.session_state["reset_grid_flag"] = True
                        st.success("¡Escalera registrada con éxito!")
                        time.sleep(0.5)
                        st.rerun()

        with tab_abm:
            st.subheader("Registro Manual de Pagos")
            col_form, col_table = st.columns([1, 2])
            with col_form:
                editing_id = st.session_state.editing_id
                df_current = get_df()
                record = df_current[df_current["id"] == editing_id].iloc[0] if editing_id is not None and not df_current[df_current["id"] == editing_id].empty else None

                st.markdown(f"##### {'✏️ Editar pago ' + editing_id if record is not None else '➕ Nuevo pago individual'}")
                proveedores_existentes = sorted(df_current["proveedor"].dropna().unique().tolist())

                with st.form("form_pago_manual", clear_on_submit=False):
                    tesoreria_sel = st.selectbox("Sede", core.TESORERIAS_VALIDAS, index=core.TESORERIAS_VALIDAS.index(record["tesoreria"]) if record is not None else 0)
                    fecha_pago = st.date_input("Fecha", value=record["fecha"].date() if record is not None and pd.notna(record["fecha"]) else dt.date.today())
                    modo_p = st.radio("Proveedor", ["Existente", "Nuevo"], horizontal=True)
                    proveedor = st.selectbox("Proveedor", proveedores_existentes, index=proveedores_existentes.index(record["proveedor"]) if record is not None else 0) if modo_p == "Existente" and proveedores_existentes else st.text_input("Nombre", value=record["proveedor"] if record is not None else "").strip().upper()
                    proyecto = st.text_input("Proyecto", value=record["proyecto"] if record is not None else "L2")
                    concepto = st.text_input("Concepto", value=record["concepto"] if record is not None else "")
                    importe = st.number_input("Importe en $ ARS", min_value=0.0, value=float(record["importe"]) if record is not None else 0.0, step=50000.0)
                    estado = st.selectbox("Estado", core.ESTADOS_VALIDOS, index=core.ESTADOS_VALIDOS.index(record["estado"]) if record is not None else 0)
                    obs = st.text_area("Notas", value=record["obs"] if record is not None else "")

                    b1, b2 = st.columns(2)
                    sub = b1.form_submit_button("💾 Guardar", use_container_width=True)
                    can = b2.form_submit_button("✖️ Cancelar", use_container_width=True)

                    if sub and proveedor:
                        df_current = get_df()
                        if record is not None:
                            idx = df_current.index[df_current["id"] == editing_id][0]
                            df_current.loc[idx, ["tesoreria", "fecha", "proveedor", "proyecto", "concepto", "moneda", "importe", "tc", "estado", "obs"]] = [tesoreria_sel, pd.Timestamp(fecha_pago), proveedor, proyecto, concepto, "ARS", importe, 1.0, estado, obs]
                            st.session_state.editing_id = None
                        else:
                            new_id = core.next_id(df_current)
                            new_row = pd.DataFrame([{"id": new_id, "tesoreria": tesoreria_sel, "fecha": pd.Timestamp(fecha_pago), "proveedor": proveedor, "proyecto": proyecto, "concepto": concepto, "moneda": "ARS", "importe": importe, "tc": 1.0, "estado": estado, "obs": obs}])
                            df_current = pd.concat([df_current, new_row], ignore_index=True)
                        set_df(df_current, sync_cloud=True)
                        st.rerun()

                    if can:
                        st.session_state.editing_id = None
                        st.rerun()

            with col_table:
                disp = df_filtered[["id", "tesoreria", "fecha", "proveedor", "concepto", "importe_ars", "estado"]].rename(columns={"id": "ID", "tesoreria": "Sede", "fecha": "Fecha", "proveedor": "Proveedor", "concepto": "Operación", "importe_ars": "Importe ARS", "estado": "Estado"}).sort_values("Fecha")
                st.dataframe(disp.style.format({"Fecha": lambda d: d.strftime("%d/%m/%Y") if pd.notna(d) else "", "Importe ARS": fmt_ars}), use_container_width=True, height=360, hide_index=True)

        with tab_export:
            st.subheader("⬇️ Exportar Base Completa")
            st.download_button("⬇️ Descargar Excel", core.export_to_excel(get_df()), f"base_pagos_{dt.date.today().isoformat()}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ===========================================================================
# MÓDULO 2: CLEARING / CHEQUES EMITIDOS (CÁMARAS BANCARIAS)
# ===========================================================================
elif modulo_activo == "🏦 Clearing / Cheques Emitidos":
    df_ch = get_cheques()

    st.title("Clearing Bancario — Cheques Emitidos")
    st.caption("Seguimiento diario de cámaras compensadoras · For Drink SA")

    # KPIs superiores proyectados estrictamente desde HOY en adelante
    kpis_ch = core.compute_clearing_kpis(df_ch, st.session_state.get("feriados", []), solo_desde_hoy=True)

    meses_con_promedios = list(kpis_ch["promedios_mensuales"].keys())
    mes_actual_str = dt.date.today().strftime("%Y-%m")
    mes_kpi_default = mes_actual_str if mes_actual_str in meses_con_promedios else (meses_con_promedios[0] if meses_con_promedios else "")

    c1, c2, c3 = st.columns([1.3, 1.2, 1.5])
    c1.metric("Total a Cubrir", fmt_ars(kpis_ch["total_comprometido"]), help="Valores activos pendientes desde hoy en adelante.")
    c2.metric("Cheques en Circulación", f"{kpis_ch['cant_cheques']} valores", help="Cantidad de cheques con vencimiento desde hoy.")

    with c3:
        if meses_con_promedios:
            c3_sel, c3_val = st.columns([1, 1.2])
            with c3_sel:
                mes_prom_sel = st.selectbox("Promedio de:", meses_con_promedios, index=meses_con_promedios.index(mes_kpi_default) if mes_kpi_default in meses_con_promedios else 0, label_visibility="collapsed")
            info_m = kpis_ch["promedios_mensuales"].get(mes_prom_sel, {})
            prom_val = info_m.get("promedio_diario", 0.0)
            dh_val = info_m.get("dias_habiles", 0)
            with c3_val:
                st.metric("Promedio Hábil Mes", fmt_ars(prom_val), help=f"Calculado sobre {dh_val} días hábiles con cheques en {mes_prom_sel}.")
        else:
            st.metric("Promedio Hábil Mes", "$ 0")

    # Micro-Badges y Barra Continua de Concentración Bancaria
    if kpis_ch.get("bancos_distribucion"):
        dist = kpis_ch["bancos_distribucion"]
        
        seg_html = "".join([
            f"<div style='flex: {info['pct']:.2f}; background-color: {BANK_THEMES.get(b, {}).get('color', '#3b82f6')}; height: 6px; border-radius: 2px;' title='{b}: {info['pct']:.1f}%'></div>"
            for b, info in dist.items()
        ])
        
        pills_html = "".join([
            f"""<div class='bank-pill' style='background:{BANK_THEMES.get(b, {}).get('bg', '#1e293b')}; border:1px solid {BANK_THEMES.get(b, {}).get('border', '#334155')}; color:#f8fafc;'>
                <span class='bank-dot' style='background-color:{BANK_THEMES.get(b, {}).get('color', '#38bdf8')};'></span>
                <b>{b}</b> {info['pct']:.1f}% <span style='color:#94a3b8; font-weight:400;'>({fmt_ars(info['monto'])})</span>
            </div>"""
            for b, info in dist.items()
        ])

        st.markdown(
            f"""
            <div style="background:#111827; border:1px solid #1e293b; border-radius:10px; padding:12px 16px; margin-top:6px; margin-bottom:14px;">
                <div style="display:flex; gap:3px; margin-bottom:10px; border-radius:3px; overflow:hidden;">
                    {seg_html}
                </div>
                <div style="display:flex; flex-wrap:wrap; gap:4px; align-items:center;">
                    {pills_html}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    tab_sabana, tab_cal, tab_carga_ch, tab_bajas = st.tabs([
        "🗓️ Sábana de Cámaras (Flujo Diario)",
        "📅 Calendario Visual",
        "➕ Carga de Cheques Emitidos",
        "🚫 Anulación de Cheques (Baja Lógica)",
    ])

    # -----------------------------------------------------------------------
    # TAB A: Sábana de Cámaras
    # -----------------------------------------------------------------------
    with tab_sabana:
        meses_unicos = sorted(df_ch["MES_KEY"].dropna().unique().tolist()) if not df_ch.empty else []
        mes_actual_default = dt.date.today().strftime("%Y-%m")
        if mes_actual_default not in meses_unicos and meses_unicos:
            mes_actual_default = meses_unicos[0]

        with st.expander("⚙️ Calibrar Semáforos por Mes y Feriados", expanded=False):
            col_m_sel, col_sem1, col_sem2, col_sem3 = st.columns([1.2, 1, 1, 1])
            with col_m_sel:
                mes_conf = st.selectbox("Mes a calibrar:", meses_unicos if meses_unicos else [mes_actual_default], key="mes_conf_sem")
            
            defaults_mes = st.session_state.semaforos_mes.get(mes_conf, (50_000_000.0, 100_000_000.0, 150_000_000.0))
            with col_sem1:
                u_verde = st.number_input("🟢 Hasta (Holgado)", min_value=1_000_000.0, value=defaults_mes[0], step=5_000_000.0, format="%.0f", key=f"uv_{mes_conf}")
            with col_sem2:
                u_amarillo = st.number_input("🟡 Hasta (Atención)", min_value=u_verde, value=defaults_mes[1], step=5_000_000.0, format="%.0f", key=f"ua_{mes_conf}")
            with col_sem3:
                u_naranja = st.number_input("🟠 Hasta (Tensión)", min_value=u_amarillo, value=defaults_mes[2], step=5_000_000.0, format="%.0f", key=f"un_{mes_conf}")

            st.session_state.semaforos_mes[mes_conf] = (u_verde, u_amarillo, u_naranja)
            st.caption(f"Semáforo guardado para **{mes_conf}**. 🔴 Por encima de {fmt_ars(u_naranja)} se destaca como Tensión Crítica.")

            st.markdown("---")
            feriados_cargados = st.session_state.get("feriados", [])
            f_feriado_nuevo = st.date_input("Agregar feriado bancario:", value=None, key="f_feriado_input")
            col_b_f1, col_b_f2 = st.columns([1, 3])
            if col_b_f1.button("Agregar feriado") and f_feriado_nuevo:
                if f_feriado_nuevo not in feriados_cargados:
                    feriados_cargados.append(f_feriado_nuevo)
                    st.session_state.feriados = sorted(feriados_cargados)
                    st.rerun()
            if feriados_cargados:
                col_b_f2.write(f"Feriados: {', '.join([d.strftime('%d/%m/%Y') for d in feriados_cargados])}")
                if st.button("Limpiar todos los feriados", key="btn_clear_feriados"):
                    st.session_state.feriados = []
                    st.rerun()

        st.divider()

        hoy_date = dt.date.today()
        col_filtro_ch, col_desc_ch, col_desc_pdf = st.columns([2, 1.2, 1.2])
        with col_filtro_ch:
            ver_desde_hoy = st.checkbox("Mostrar únicamente desde hoy en adelante", value=False)
            f_desde_limite = hoy_date if ver_desde_hoy else None

        matriz_ch, bancos_activos = core.build_clearing_matrix(df_ch, fecha_inicio=f_desde_limite)

        if matriz_ch.empty:
            st.info("No hay cheques pendientes registrados para el período seleccionado.")
        else:
            def style_clearing(row):
                is_sub = row.get("IS_SUBTOTAL", False)
                styles = [""] * len(row)

                if is_sub:
                    sub_style = "background-color: #1e293b; font-weight: 700; color: #38bdf8; border-top: 1px solid #334155; border-bottom: 2px solid #475569;"
                    return [sub_style] * len(row)

                val = row["TOTAL"]
                mes_k = row.get("MES_KEY", "")
                u_v, u_a, u_n = st.session_state.semaforos_mes.get(mes_k, (50_000_000.0, 100_000_000.0, 150_000_000.0))

                if val > u_n:
                    style_tot = "background-color: #7f1d1d; color: #fecaca; font-weight: 700;"
                elif val > u_a:
                    style_tot = "background-color: #7c2d12; color: #ffedd5; font-weight: 700;"
                elif val > u_v:
                    style_tot = "background-color: #78350f; color: #fef3c7; font-weight: 700;"
                else:
                    style_tot = "background-color: #064e3b; color: #d1fae5; font-weight: 700;"

                idx_tot = list(row.index).index("TOTAL")
                styles[idx_tot] = style_tot
                return styles

            def export_clearing_excel(pivot_df: pd.DataFrame, bancos_list: list[str]) -> bytes:
                wb = Workbook()
                ws = wb.active
                ws.title = "Camaras"

                ws["A1"] = "CÁMARAS COMPENSADORAS — CHEQUES EMITIDOS"
                ws["A1"].font = Font(size=13, bold=True, color="1F4E78")
                ws["A2"] = f"Generado el: {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}"
                ws["A2"].font = Font(size=9, italic=True, color="666666")

                headers = ["FECHA"] + bancos_list + ["TOTAL"]
                h_row = 4
                for c_idx, h in enumerate(headers, 1):
                    cell = ws.cell(row=h_row, column=c_idx, value=h)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                for r_idx, (_, r) in enumerate(pivot_df.iterrows()):
                    row_num = h_row + 1 + r_idx
                    is_sub = r.get("IS_SUBTOTAL", False)
                    ws.cell(row=row_num, column=1, value=r["FECHA_LABEL"]).alignment = Alignment(horizontal="center" if not is_sub else "left")

                    for b_idx, b in enumerate(bancos_list, 2):
                        v = float(r.get(b, 0.0))
                        cell_b = ws.cell(row=row_num, column=b_idx, value=v if v > 0 else "")
                        if v > 0:
                            cell_b.number_format = "$ #,##0"

                    tot_v = float(r["TOTAL"])
                    cell_t = ws.cell(row=row_num, column=len(headers), value=tot_v)
                    cell_t.font = Font(bold=True)
                    cell_t.number_format = "$ #,##0"

                    if is_sub:
                        for col_idx in range(1, len(headers) + 1):
                            c_sub = ws.cell(row=row_num, column=col_idx)
                            c_sub.font = Font(bold=True)
                            c_sub.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

                for col in ws.columns:
                    max_len = max(len(str(cell.value or "")) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

                buf = BytesIO()
                wb.save(buf)
                return buf.getvalue()

            def export_clearing_pdf(pivot_df: pd.DataFrame, bancos_list: list[str]) -> bytes:
                buf = BytesIO()
                doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=25, bottomMargin=25)
                styles = getSampleStyleSheet()
                story = []

                title_style = ParagraphStyle('ClrTitle', parent=styles['Heading1'], fontSize=13, textColor=colors.HexColor('#1F4E78'), spaceAfter=4)
                sub_style = ParagraphStyle('ClrSub', parent=styles['Normal'], fontSize=8.5, textColor=colors.HexColor('#555555'), spaceAfter=10)

                story.append(Paragraph("<b>FOR DRINK SA — SÁBANA DE CÁMARAS COMPENSADORAS</b>", title_style))
                story.append(Paragraph(f"Emisión: {dt.datetime.now().strftime('%d/%m/%Y %H:%M')} | Valores pendientes", sub_style))

                headers = ["FECHA"] + bancos_list + ["TOTAL"]
                table_data = [headers]

                for _, r in pivot_df.iterrows():
                    row_vals = [r["FECHA_LABEL"]]
                    for b in bancos_list:
                        v = float(r.get(b, 0.0))
                        row_vals.append(fmt_ars(v) if v > 0 else "")
                    row_vals.append(fmt_ars(r["TOTAL"]))
                    table_data.append(row_vals)

                col_w = max(42, int(780 / len(headers)))
                t = Table(table_data, colWidths=[col_w] * len(headers))
                t_styles = [
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 7.5),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#DDDDDD')),
                    ('FONTSIZE', (0, 1), (-1, -1), 7),
                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ]
                for r_i, (_, r) in enumerate(pivot_df.iterrows(), 1):
                    if r.get("IS_SUBTOTAL"):
                        t_styles.append(('BACKGROUND', (0, r_i), (-1, r_i), colors.HexColor('#E2E8F0')))
                        t_styles.append(('FONTNAME', (0, r_i), (-1, r_i), 'Helvetica-Bold'))

                t.setStyle(TableStyle(t_styles))
                story.append(t)
                doc.build(story)
                return buf.getvalue()

            with col_desc_ch:
                st.download_button(
                    "📥 Descargar Excel",
                    export_clearing_excel(matriz_ch, bancos_activos),
                    f"Clearing_{dt.date.today().strftime('%Y%m%d')}.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

            with col_desc_pdf:
                st.download_button(
                    "📄 Descargar PDF",
                    export_clearing_pdf(matriz_ch, bancos_activos),
                    f"Clearing_{dt.date.today().strftime('%Y%m%d')}.pdf",
                    "application/pdf",
                    use_container_width=True,
                )

            view_matriz = matriz_ch.rename(columns={"FECHA_LABEL": "FECHA"}).drop(columns=["Fecha Pago", "PERIODO"], errors="ignore")
            format_dict = {b: lambda v: fmt_ars(v) if v > 0 else "—" for b in bancos_activos}
            format_dict["TOTAL"] = fmt_ars

            cols_to_show = ["FECHA"] + bancos_activos + ["TOTAL", "IS_SUBTOTAL", "MES_KEY"]
            st.dataframe(
                view_matriz[cols_to_show].style.apply(style_clearing, axis=1).format(format_dict),
                use_container_width=True,
                height=580,
                hide_index=True,
                column_config={
                    "IS_SUBTOTAL": None,
                    "MES_KEY": None,
                },
            )

    # -----------------------------------------------------------------------
    # TAB B: Calendario Visual Clásico y Estable (Sin Deformaciones)
    # -----------------------------------------------------------------------
    with tab_cal:
        col_cal_m, col_cal_y = st.columns([1, 1])
        mes_cal_sel = col_cal_m.selectbox("Mes", list(core.MESES_ES.values()), index=dt.date.today().month - 1, key="cal_mes_select")
        ano_cal_sel = col_cal_y.number_input("Año", min_value=2024, max_value=2030, value=dt.date.today().year, key="cal_ano_input")

        num_mes_sel = [k for k, v in core.MESES_ES.items() if v == mes_cal_sel][0]
        cal = calendar.monthcalendar(ano_cal_sel, num_mes_sel)
        dias_nombres = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

        c_cols = st.columns(7)
        for i, nom_d in enumerate(dias_nombres):
            c_cols[i].markdown(f"<div style='text-align:center; font-weight:700; font-size:12px; color:#94a3b8; padding-bottom:6px;'>{nom_d}</div>", unsafe_allow_html=True)

        df_activos_cal = df_ch[df_ch.get("Estado", "Emitido").astype(str).str.lower() != "anulado"]

        for week in cal:
            w_cols = st.columns(7)
            for d_idx, day in enumerate(week):
                with w_cols[d_idx]:
                    if day == 0:
                        st.markdown("<div style='min-height:96px;'></div>", unsafe_allow_html=True)
                    else:
                        fecha_d = dt.date(ano_cal_sel, num_mes_sel, day)
                        df_dia = df_activos_cal[df_activos_cal["Fecha Pago"].dt.date == fecha_d]
                        tot_dia = df_dia["Importe"].sum() if not df_dia.empty else 0.0

                        es_feriado = fecha_d in st.session_state.get("feriados", [])
                        lbl_feriado = " <span style='color:#f59e0b; font-size:10px;'>(Feriado)</span>" if es_feriado else ""

                        if tot_dia > 0:
                            items_bancos = ""
                            por_banco = df_dia.groupby("Banco")["Importe"].sum()
                            for b, m in por_banco.items():
                                items_bancos += f"<div style='font-size:10px; color:#cbd5e1; margin-bottom:2px;'>• {b}: <b>{fmt_ars(m)}</b></div>"

                            tot_html = f"<div class='cal-day-total'>{fmt_ars(tot_dia)}</div>"
                            details_html = f"""
                            <details style="margin-top:4px;">
                                <summary class="cal-details-summary">▾ Ver detalle</summary>
                                <div style="margin-top:4px; padding-top:4px; border-top:1px solid #334155;">
                                    {items_bancos}
                                </div>
                            </details>
                            """
                        else:
                            tot_html = "<div style='font-size:11px; color:#64748b; margin-top:6px;'>Sin vencimientos</div>"
                            details_html = ""

                        st.markdown(
                            f"""
                            <div class='cal-day-box'>
                                <div>
                                    <div class='cal-day-header'>
                                        <span>{day:02d}</span>{lbl_feriado}
                                    </div>
                                    {tot_html}
                                </div>
                                {details_html}
                            </div>
                            """,
                            unsafe_allow_html=True,
                        )

    # -----------------------------------------------------------------------
    # TAB C: Carga Masiva e Individual (Blindada contra Duplicados y Errores)
    # -----------------------------------------------------------------------
    with tab_carga_ch:
        st.subheader("➕ Cargar Nuevos Cheques Emitidos")
        st.caption("Validación automática de duplicados y blindaje de importación desde Excel.")

        if "cheque_guardado_msj" in st.session_state:
            st.success(st.session_state.cheque_guardado_msj, icon="✅")
            del st.session_state["cheque_guardado_msj"]

        modo_carga_ch = st.radio(
            "Modalidad de carga:",
            ["📋 Pegar desde Excel (Masivo)", "✍️ Carga individual"],
            horizontal=True,
            key="modo_carga_ch_radio",
        )

        bancos_predefinidos = ["BBVA", "BNA", "GALICIA", "MACRO", "PROVINCIA", "SANTANDER", "SUPERVIELLE"]

        if modo_carga_ch == "📋 Pegar desde Excel (Masivo)":
            st.caption("Copiá las columnas: `Banco` | `Fecha Emisión` | `Fecha Pago` | `Nro. Cheque` | `Importe` | `Beneficiario` (con o sin encabezados):")
            txt_ch_excel = st.text_area(
                "Pegar celdas copiadas",
                placeholder="GALICIA\t24/09/2026\t21/10/2026\t1433\t$ 5.144.292,18\tFAMIQ SRL\nSANTANDER\t24/09/2026\t29/10/2026\t2334\t8.765.304,12\tESTABLECIMIENTO GRAFICO",
                height=130,
                key="txt_ch_excel_input",
            )

            if st.button("📥 Procesar y Guardar Cheques en Google Sheets"):
                if not txt_ch_excel.strip():
                    st.warning("El campo de texto está vacío.")
                else:
                    lines = [l.strip() for l in txt_ch_excel.strip().split("\n") if l.strip()]
                    nuevos_cheques = []
                    for line in lines:
                        p = re.split(r"[\t;|]+", line)
                        if len(p) >= 5:
                            b_val = p[0].strip().upper()
                            # Blindaje: omitir si es la fila de encabezados de Excel
                            if b_val in ("BANCO", "BANK", "ENTIDAD") or "CHEQUE" in str(p[3]).upper():
                                continue

                            fe_val = pd.to_datetime(p[1].strip(), dayfirst=True, errors="coerce")
                            fp_val = pd.to_datetime(p[2].strip(), dayfirst=True, errors="coerce")
                            nro_val = str(p[3]).strip()
                            
                            # Conversión numérica protegida contra ValueError
                            imp_raw = str(p[4]).replace("$", "").replace(" ", "").replace(".", "").replace(",", ".").strip()
                            try:
                                imp_val = float(imp_raw) if imp_raw else 0.0
                            except (ValueError, TypeError):
                                continue

                            benef_val = p[5].strip() if len(p) > 5 else ""

                            if pd.notna(fp_val) and imp_val > 0:
                                nuevos_cheques.append({
                                    "Banco": b_val,
                                    "EMPRESA": "FD",
                                    "Cuenta Libradora": "",
                                    "Fecha Emisión": fe_val if pd.notna(fe_val) else fp_val,
                                    "Fecha Pago": fp_val,
                                    "Nro. de Cheque": nro_val,
                                    "Importe": imp_val,
                                    "CUIT Beneficiario": "",
                                    "Razón Social Beneficiario": benef_val,
                                    "Estado": "Emitido",
                                    "Motivo Anulación": "",
                                })

                    if not nuevos_cheques:
                        st.error("No se encontraron filas válidas para procesar. Revisá los datos copiados.")
                    else:
                        # Validación de duplicados
                        df_actual = get_cheques()
                        validos, duplicados = core.check_cheque_duplicates(df_actual, nuevos_cheques)

                        if duplicados:
                            det_dups = ", ".join([f"#{d['Nro. de Cheque']} ({d['Banco']})" for d in duplicados[:5]])
                            st.warning(f"⚠️ Se detectaron **{len(duplicados)} cheques duplicados** ya existentes en la base y fueron omitidos: {det_dups}")

                        if validos:
                            df_nuevos = pd.DataFrame(validos)
                            df_act = pd.concat([df_actual, df_nuevos], ignore_index=True)
                            if save_cheques_source(df_act):
                                st.session_state.df_cheques = core.normalize_cheques_df(df_act)
                                st.session_state.cheque_guardado_msj = f"¡Se guardaron exitosamente {len(validos)} cheques en la base de datos!"
                                st.rerun()
                        else:
                            st.error("Todos los cheques pegados ya figuraban previamente en la base de datos.")
        else:
            with st.form("form_nuevo_cheque"):
                c_ch1, c_ch2, c_ch3 = st.columns(3)
                with c_ch1:
                    ch_banco = st.selectbox("Banco Emisor", bancos_predefinidos)
                    ch_empresa = st.text_input("Empresa", value="FD")
                    ch_cuenta = st.text_input("Cuenta Libradora", placeholder="Ej: 16603")
                with c_ch2:
                    ch_fe = st.date_input("Fecha Emisión", value=dt.date.today())
                    ch_fp = st.date_input("Fecha Pago", value=dt.date.today() + dt.timedelta(days=7))
                    ch_nro = st.text_input("Nro. de Cheque", placeholder="Ej: 1264")
                with c_ch3:
                    ch_imp = st.number_input("Importe en $ ARS", min_value=1.0, step=50000.0, format="%.2f")
                    ch_cuit_b = st.text_input("CUIT Beneficiario", placeholder="Opcional")
                    ch_razon_b = st.text_input("Razón Social Beneficiario", placeholder="Ej: HAASEN SA")

                if st.form_submit_button("💾 Guardar Cheque", use_container_width=True):
                    df_actual = get_cheques()
                    item_single = [{
                        "Banco": ch_banco,
                        "EMPRESA": ch_empresa,
                        "Cuenta Libradora": ch_cuenta,
                        "Fecha Emisión": pd.Timestamp(ch_fe),
                        "Fecha Pago": pd.Timestamp(ch_fp),
                        "Nro. de Cheque": str(ch_nro).strip(),
                        "Importe": float(ch_imp),
                        "CUIT Beneficiario": ch_cuit_b,
                        "Razón Social Beneficiario": ch_razon_b,
                        "Estado": "Emitido",
                        "Motivo Anulación": "",
                    }]
                    validos, duplicados = core.check_cheque_duplicates(df_actual, item_single)
                    if duplicados:
                        st.error(f"❌ Error de duplicidad: El cheque N° {ch_nro} para {ch_banco} ya existe en la base de datos.")
                    else:
                        df_act = pd.concat([df_actual, pd.DataFrame(validos)], ignore_index=True)
                        if save_cheques_source(df_act):
                            st.session_state.df_cheques = core.normalize_cheques_df(df_act)
                            st.session_state.cheque_guardado_msj = f"¡Cheque N° {ch_nro} guardado exitosamente!"
                            st.rerun()

    # -----------------------------------------------------------------------
    # TAB D: Anulación de Cheques (Baja Lógica / Trazabilidad)
    # -----------------------------------------------------------------------
    with tab_bajas:
        st.subheader("🚫 Anulación de Cheques (Baja Lógica)")
        st.caption("Anular un cheque lo excluye inmediatamente del clearing y del calendario, preservando el registro histórico en Google Sheets.")

        activos = df_ch[df_ch.get("Estado", "Emitido").astype(str).str.lower() != "anulado"].copy()
        if activos.empty:
            st.info("No hay cheques activos en circulación para anular.")
        else:
            col_b_an1, col_b_an2 = st.columns([1, 1])
            with col_b_an1:
                bancos_con_ch = sorted(activos["Banco"].unique().tolist())
                b_filtro_an = st.selectbox("Filtrar por Banco:", ["Todos"] + bancos_con_ch, key="sel_b_anular")
            
            df_para_anular = activos if b_filtro_an == "Todos" else activos[activos["Banco"] == b_filtro_an]

            opciones_anular = [
                f"#{r['Nro. de Cheque']} | {r['Banco']} | {fmt_ars(r['Importe'])} | {r.get('Razón Social Beneficiario', '')} (Vto: {r['Fecha Pago'].strftime('%d/%m/%Y') if pd.notna(r['Fecha Pago']) else ''})"
                for _, r in df_para_anular.iterrows()
            ]

            ch_sel_anular = st.selectbox("Seleccionar Cheque a Anular:", opciones_anular, key="sel_ch_anular_box")
            motivo_anulacion = st.text_input("Motivo de anulación (opcional):", placeholder="Ej: Error en fecha / Reemplazado por transferencia", key="motivo_anular_txt")

            if st.button("🚫 Proceder con la Anulación del Cheque"):
                nro_extraido = ch_sel_anular.split("|")[0].replace("#", "").strip()
                banco_extraido = ch_sel_anular.split("|")[1].strip()

                df_mod = get_cheques().copy()
                mask = (df_mod["Nro. de Cheque"].astype(str).str.strip() == nro_extraido) & (df_mod["Banco"].astype(str).str.strip().str.upper() == banco_extraido.upper())
                
                if mask.any():
                    df_mod.loc[mask, "Estado"] = "Anulado"
                    df_mod.loc[mask, "Motivo Anulación"] = motivo_anulacion or "Anulado por Tesorería"
                    if save_cheques_source(df_mod):
                        st.session_state.df_cheques = core.normalize_cheques_df(df_mod)
                        st.success(f"¡El cheque N° {nro_extraido} de {banco_extraido} ha sido ANULADO exitosamente y excluido del Clearing!")
                        time.sleep(0.5)
                        st.rerun()

# ===========================================================================
# MÓDULO 3: ANALIZADOR DE LIBRADORES · BCRA (NATIVO STREAMLIT DARK)
# ===========================================================================
elif modulo_activo == "🔍 Analizador de Libradores · BCRA":
    col_t_bcra, col_btns_bcra = st.columns([3, 1.2])
    with col_t_bcra:
        st.title("Analizador de Libradores BCRA")
        st.caption("Evaluación crediticia de CUITs en Central de Deudores y Cheques Rechazados en tiempo real.")

    with col_btns_bcra:
        st.write("")
        b_demo, b_clean = st.columns(2)
        if b_demo.button("Demo", use_container_width=True):
            st.session_state["bcra_input_val"] = "30719231116\n30711743711\n30719168295\n20222645019\n20165261624\n30717693392"
            st.rerun()
        if b_clean.button("Limpiar", use_container_width=True):
            st.session_state["bcra_input_val"] = ""
            st.session_state["resultados_bcra"] = []
            st.session_state["bcra_status_msg"] = None
            st.rerun()

    input_text_val = st.session_state.get("bcra_input_val", "")

    c_box_in, c_box_btn = st.columns([3.5, 1])
    with c_box_in:
        cuits_raw = st.text_area(
            "CUITs a evaluar",
            value=input_text_val,
            placeholder="Pegá los CUITs (uno por línea o separados por espacio/coma)\nEj:\n30-71649553-8\n30-52801149-3",
            height=95,
            label_visibility="collapsed",
            key="area_cuits_bcra",
        )
    with c_box_btn:
        st.write("")
        st.write("")
        btn_consultar_bcra = st.button("🔎 Consultar BCRA", use_container_width=True)

    if btn_consultar_bcra and cuits_raw.strip():
        lista_cuits = core.parse_cuits_input(cuits_raw)
        if not lista_cuits:
            st.warning("No se detectaron CUITs válidos de 11 dígitos.")
        else:
            prog_bar = st.progress(0, text="Iniciando consultas al BCRA...")
            res_bcra = []
            for idx, c in enumerate(lista_cuits):
                prog_bar.progress((idx + 1) / len(lista_cuits), text=f"Consultando CUIT {c} ({idx + 1}/{len(lista_cuits)})...")
                res = core.fetch_bcra_data(c)
                res_bcra.append(res)
                time.sleep(0.2)
            prog_bar.empty()

            peso = {"bad": 0, "warn": 1, "ok": 2}
            res_bcra.sort(key=lambda x: peso.get(x["risk"], 3))
            st.session_state["resultados_bcra"] = res_bcra
            st.session_state["bcra_status_msg"] = f"✅ Consulta finalizada ({len(res_bcra)} libradores evaluados)."

    if st.session_state.get("bcra_status_msg"):
        st.caption(st.session_state["bcra_status_msg"])

    data_bcra = st.session_state.get("resultados_bcra", [])
    n_tot = len(data_bcra)
    n_ok = len([x for x in data_bcra if x["risk"] == "ok"])
    n_warn = len([x for x in data_bcra if x["risk"] == "warn"])
    n_bad = len([x for x in data_bcra if x["risk"] == "bad"])

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.markdown(f"<div class='bcra-card-dark'><div class='bcra-label-dark'>Libradores Evaluados</div><div class='bcra-val-dark' style='color:#f8fafc;'>{n_tot}</div></div>", unsafe_allow_html=True)
    with c2:
        st.markdown(f"<div class='bcra-card-dark'><div class='bcra-label-dark'>Sin Alertas</div><div class='bcra-val-dark' style='color:#10b981;'>{n_ok}</div></div>", unsafe_allow_html=True)
    with c3:
        st.markdown(f"<div class='bcra-card-dark'><div class='bcra-label-dark'>Revisar</div><div class='bcra-val-dark' style='color:#f59e0b;'>{n_warn}</div></div>", unsafe_allow_html=True)
    with c4:
        st.markdown(f"<div class='bcra-card-dark'><div class='bcra-label-dark'>Alertas / Rechazar</div><div class='bcra-val-dark' style='color:#ef4444;'>{n_bad}</div></div>", unsafe_allow_html=True)

    st.write("")

    if data_bcra:
        col_grid_main, col_grid_drawer = st.columns([1.6, 1.1])

        with col_grid_main:
            st.markdown(
                """
                <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:8px;">
                    <b style="font-size:14px; color:#f8fafc;">Listado de Libradores</b>
                    <small style="color:#94a3b8;">Prioridad: Alerta → Revisar → Sin Alertas</small>
                </div>
                """,
                unsafe_allow_html=True,
            )

            table_rows = []
            for x in data_bcra:
                sit = x["worst"]
                if sit in (0, 1):
                    badge_style = "background:#064e3b; color:#34d399; border: 1px solid #059669;"
                elif sit == 2:
                    badge_style = "background:#78350f; color:#fbbf24; border: 1px solid #d97706;"
                else:
                    badge_style = "background:#7f1d1d; color:#f87171; border: 1px solid #dc2626;"

                if x["risk"] == "bad":
                    pill_style = "background:#7f1d1d; color:#fecaca; border:1px solid #dc2626;"
                elif x["risk"] == "warn":
                    pill_style = "background:#78350f; color:#fde68a; border:1px solid #d97706;"
                else:
                    pill_style = "background:#064e3b; color:#a7f3d0; border:1px solid #059669;"

                imp_txt = f"<span style='color:#ef4444; font-weight:700;'>({x['pending']} impagos)</span>" if x["pending"] > 0 else "<span style='color:#94a3b8;'>(0 impagos)</span>"

                table_rows.append(
                    f"<tr style='border-bottom: 1px solid #1e293b;'>"
                    f"<td style='padding:10px 8px; color:#f8fafc;'><b>{x['denominacion']}</b><br><small style='font-family:monospace; color:#94a3b8;'>{x['cuit']}</small></td>"
                    f"<td style='text-align:center; padding:10px 4px;'><span style='display:inline-flex; align-items:center; justify-content:center; width:24px; height:24px; border-radius:50%; font-weight:700; font-size:11.5px; {badge_style}'>{sit}</span></td>"
                    f"<td style='padding:10px 8px; font-weight:700; color:#f8fafc;'>{fmt_ars(x['debt'])}</td>"
                    f"<td style='padding:10px 8px; color:#f8fafc;'>{x['rejected']} {imp_txt}</td>"
                    f"<td style='padding:10px 8px; text-align:right;'><span style='display:inline-block; padding:3px 9px; border-radius:9999px; font-size:11px; font-weight:700; {pill_style}'>{x['risk_label']}</span></td>"
                    f"</tr>"
                )

            html_table = (
                f"<div class='bcra-table-box'>"
                f"<table style='width:100%; border-collapse:collapse; font-size:13px;'>"
                f"<thead><tr style='border-bottom:1px solid #334155; color:#94a3b8; font-size:10.5px; text-transform:uppercase;'>"
                f"<th style='padding:8px; text-align:left;'>Librador / Denominación</th>"
                f"<th style='padding:8px; text-align:center;'>Peor Sit.</th>"
                f"<th style='padding:8px; text-align:left;'>Deuda Bancaria</th>"
                f"<th style='padding:8px; text-align:left;'>Rechazos (Pend.)</th>"
                f"<th style='padding:8px; text-align:right;'>Estado</th>"
                f"</tr></thead>"
                f"<tbody>{''.join(table_rows)}</tbody>"
                f"</table></div>"
            )
            st.markdown(html_table, unsafe_allow_html=True)

        with col_grid_drawer:
            st.markdown(
                """
                <div style="margin-bottom:8px;">
                    <span class="bcra-label-dark">Ficha del Librador</span>
                </div>
                """,
                unsafe_allow_html=True,
            )

            opciones_nombres = [f"{x['denominacion']} ({x['cuit']})" for x in data_bcra]
            sel_lib = st.selectbox("Seleccionar:", opciones_nombres, label_visibility="collapsed", key="sel_drawer_lib")
            idx_sel = opciones_nombres.index(sel_lib)
            lib = data_bcra[idx_sel]

            if lib["risk"] == "bad":
                badge_d_style = "background:#7f1d1d; color:#fecaca; border:1px solid #dc2626;"
            elif lib["risk"] == "warn":
                badge_d_style = "background:#78350f; color:#fde68a; border:1px solid #d97706;"
            else:
                badge_d_style = "background:#064e3b; color:#a7f3d0; border:1px solid #059669;"

            drawer_html = (
                f"<div class='bcra-drawer-box'>"
                f"<h3 style='margin:0 0 2px 0; font-size:16px; font-weight:700; color:#f8fafc;'>{lib['denominacion']}</h3>"
                f"<div style='font-family:monospace; color:#94a3b8; font-size:12px; margin-bottom:10px;'>{lib['cuit']}</div>"
                f"<span style='display:inline-block; padding:3px 9px; border-radius:9999px; font-size:11px; font-weight:700; {badge_d_style}'>{lib['risk_label']}</span>"
                f"<div style='display:grid; grid-template-columns:1fr 1fr; gap:8px; margin-top:14px;'>"
                f"<div class='kpi-mini-box'><div class='bcra-label-dark'>Peor Situación</div><b style='font-size:15px; color:#f8fafc; margin-top:2px; display:block;'>{lib['worst']}</b></div>"
                f"<div class='kpi-mini-box'><div class='bcra-label-dark'>Deuda Total</div><b style='font-size:15px; color:#f8fafc; margin-top:2px; display:block;'>{fmt_ars(lib['debt'])}</b></div>"
                f"<div class='kpi-mini-box'><div class='bcra-label-dark'>Total Rechazos</div><b style='font-size:15px; color:#f8fafc; margin-top:2px; display:block;'>{lib['rejected']}</b></div>"
                f"<div class='kpi-mini-box'><div class='bcra-label-dark'>Impagos / Pend.</div><b style='font-size:15px; margin-top:2px; display:block; color:{'#ef4444' if lib['pending'] > 0 else '#f8fafc'};'>{lib['pending']}</b></div>"
                f"</div></div>"
            )
            st.markdown(drawer_html, unsafe_allow_html=True)

            st.write("")
            for al in lib["alerts"]:
                if lib["risk"] == "bad":
                    al_style = "background:#450a0a; color:#fca5a5; border-left: 3px solid #ef4444;"
                elif lib["risk"] == "warn":
                    al_style = "background:#451a03; color:#fcd34d; border-left: 3px solid #f59e0b;"
                else:
                    al_style = "background:#064e3b; color:#6ee7b7; border-left: 3px solid #10b981;"
                st.markdown(f"<div style='padding:7px 10px; border-radius:6px; font-size:12px; font-weight:500; margin-bottom:5px; {al_style}'>{al}</div>", unsafe_allow_html=True)

            if lib["last3"]:
                st.markdown("<small style='font-weight:600; text-transform:uppercase; color:#94a3b8;'>Últimos 3 Cheques Rechazados</small>", unsafe_allow_html=True)
                for ch in lib["last3"]:
                    f_ch = str(ch.get("fechaRechazo", "—")).split("T")[0]
                    if "-" in f_ch:
                        p_f = f_ch.split("-")
                        if len(p_f) == 3 and len(p_f[0]) == 4:
                            f_ch = f"{p_f[2]}-{p_f[1]}-{p_f[0]}"
                    m_ch = fmt_ars(ch.get("monto", 0))
                    estado_ch = "<span style='color:#34d399; font-weight:700;'>Pagado</span>" if ch.get("fechaPago") else "<span style='color:#f87171; font-weight:700;'>Impago</span>"
                    st.markdown(f"<div style='font-size:12px; border-bottom:1px solid #1e293b; padding:4px 0; color:#cbd5e1;'>• <b>{f_ch}</b> — {m_ch} ({ch.get('causal')}) [{estado_ch}]</div>", unsafe_allow_html=True)

        st.divider()
        st.subheader("📱 Mensaje para WhatsApp")
        rechazados = [x for x in data_bcra if x["risk"] == "bad"]
        revisar = [x for x in data_bcra if x["risk"] == "warn"]

        if rechazados:
            wa_txt = "Hola! Te paso el resultado de la tanda analizada:\n\n"
            wa_txt += "❌ *RECHAZAR los siguientes libradores:*\n"
            for r in rechazados:
                motivo = f"Sit: {r['worst']}"
                if r['pending'] > 0:
                    motivo += f" | {r['pending']} cheques impagos"
                if r['rejected'] > 5:
                    motivo += f" | {r['rejected']} rechazos totales"
                wa_txt += f"• *{r['denominacion']}* (CUIT {r['cuit']}) - {motivo}\n"
            if revisar:
                wa_txt += "\n⚠️ *A REVISAR antes de recibir:*\n"
                for rv in revisar:
                    wa_txt += f"• {rv['denominacion']} (Sit: {rv['worst']})\n"
        elif revisar:
            wa_txt = "Hola! De la tanda analizada no hay alertas críticas, pero sugiero *revisar*:\n\n"
            for rv in revisar:
                wa_txt += f"• *{rv['denominacion']}* (CUIT {rv['cuit']}) - Sit: {rv['worst']}\n"
        else:
            wa_txt = "Hola! Todos los libradores analizados están en condiciones *OK (Sin Alertas)*. ✅"

        st.text_area("Copia el texto directamente para enviar:", value=wa_txt, height=120)
